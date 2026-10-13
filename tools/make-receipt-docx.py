#!/usr/bin/env python3
"""領収書1枚ぶんの .docx を作る。

    python3 tools/make-receipt-docx.py '住田　様' '2026年11月15日' out.docx

Googleドライブにアップロードすると、そのままGoogleドキュメントになる。
HTMLからの変換と違って画像が確実に入るので、こちらを使う。
"""
import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

IMG = pathlib.Path(__file__).resolve().parent.parent / 'assets' / 'docx'
JP = 'Noto Sans JP'

INK = RGBColor(0x20, 0x21, 0x24)
TITLE = RGBColor(0x3B, 0x3F, 0x5C)
GREY = RGBColor(0x5F, 0x63, 0x68)
RULE = RGBColor(0x99, 0x99, 0x99)


def run(p, text, size=10, bold=False, color=INK):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = JP
    r._element.rPr.rFonts.set(qn('w:eastAsia'), JP)
    return r


def para(doc, align=None, before=0, after=0, line=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if line:
        p.paragraph_format.line_spacing = line
    return p


def shade(cell, hex_color):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, size=10, bold=False, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run(p, text, size=size, bold=bold)


def table_width(table, widths):
    """列幅を実際に効かせる。tblGrid と各セルの両方を指定し、
    レイアウトを fixed にしないと Word / Google が勝手に割り振ってしまう。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), str(int(sum(widths) * 1440)))
    w.set(qn('w:type'), 'dxa')
    tblPr.append(w)

    grid = tbl.find(qn('w:tblGrid'))
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(width * 1440)))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def bottom_border(p, hex_color='C9C6C1'):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), hex_color)
    bdr.append(bottom)
    pPr.append(bdr)


def build(to_name, issue_date, out_path):
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Inches(0.28)
    s.bottom_margin = Inches(0.28)
    s.left_margin = Inches(0.69)
    s.right_margin = Inches(0.69)

    style = doc.styles['Normal']
    style.font.name = JP
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), JP)

    # ヘッダーのイラスト
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(IMG / 'receipt-header.png'), width=Inches(1.85))

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    run(p, '♪　領　収　書　♪', size=18, bold=True, color=TITLE)

    p = para(doc, after=10)
    bottom_border(p)

    p = para(doc, WD_ALIGN_PARAGRAPH.RIGHT, after=10)
    run(p, f'発行日：{issue_date}', size=10, color=GREY)

    p = para(doc, after=1)
    run(p, to_name, size=16, bold=True)
    p = para(doc, after=9)
    run(p, '─' * 20, size=10, color=RULE)

    p = para(doc, after=9)
    run(p, '件名：1曲ひいてみよう やってみようコンサート　参加費として', size=10.5)

    p = para(doc, after=3)
    run(p, '合計金額　　¥8,800-（税込）', size=17, bold=True)
    p = para(doc, after=10)
    run(p, '上記の金額を正に領収いたしました。', size=10)

    # 品目の表
    t = doc.add_table(rows=2, cols=4)
    t.style = 'Table Grid'
    t.autofit = False
    table_width(t, (3.5, 0.9, 1.3, 1.3))
    for cell, text in zip(t.rows[0].cells, ('品目', '数量', '単価', '金額')):
        shade(cell, 'EEEDFC')
        cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    body = ('参加費（1曲ひいてみよう やってみようコンサート）', '1', '¥8,000', '¥8,000')
    aligns = (None, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT)
    for cell, text, al in zip(t.rows[1].cells, body, aligns):
        cell_text(cell, text, align=al)

    para(doc, after=5)

    # 小計・消費税・合計
    t2 = doc.add_table(rows=3, cols=2)
    t2.style = 'Table Grid'
    t2.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table_width(t2, (1.6, 1.3))
    for r, (label, value) in enumerate((('小計', '¥8,000'), ('消費税（10%）', '¥800'), ('合計', '¥8,800'))):
        cell_text(t2.rows[r].cells[0], label)
        cell_text(t2.rows[r].cells[1], value, align=WD_ALIGN_PARAGRAPH.RIGHT)

    para(doc, after=6)

    p = para(doc, after=1)
    run(p, '発行者', size=10, bold=True)
    p = para(doc, after=1)
    run(p, '増見麻美　', size=10)
    p.add_run().add_picture(str(IMG / 'seal.png'), width=Inches(0.64))
    p = para(doc, after=1)
    run(p, '〒456-0062　愛知県名古屋市熱田区大宝４−１５−２１', size=10)
    p = para(doc, after=10)
    run(p, '登録番号：T2810958830122', size=10)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    run(p, 'ご参加いただき、ありがとうございます。', size=9.5, color=GREY)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(str(IMG / 'receipt-footer.png'), width=Inches(3.4))

    doc.save(out_path)
    slim(out_path)
    return out_path


MINIMAL_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:docDefaults><w:rPrDefault><w:rPr>
<w:rFonts w:ascii="Noto Sans JP" w:hAnsi="Noto Sans JP" w:eastAsia="Noto Sans JP" w:cs="Noto Sans JP"/>
<w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr></w:rPrDefault>
<w:pPrDefault><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr></w:pPrDefault>
</w:docDefaults>
<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>
<w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/>
<w:tblPr><w:tblBorders>
<w:top w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
<w:left w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
<w:right w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
<w:insideV w:val="single" w:sz="4" w:space="0" w:color="D4D1CC"/>
</w:tblBorders></w:tblPr></w:style>
</w:styles>"""


def slim(path):
    """雛形についてくる使わないパーツを落として軽くする。
    アップロードのたびに丸ごと送ることになるので、小さいほどよい。"""
    import re
    import zipfile

    drop = ('word/stylesWithEffects.xml', 'docProps/thumbnail.jpeg',
            'customXml/item1.xml', 'customXml/itemProps1.xml',
            'customXml/_rels/item1.xml.rels', 'word/numbering.xml',
            'word/theme/theme1.xml', 'word/webSettings.xml', 'docProps/app.xml',
            'word/settings.xml', 'word/fontTable.xml', 'docProps/core.xml')
    src = zipfile.ZipFile(path)
    items = {i.filename: src.read(i.filename) for i in src.infolist()}
    src.close()

    for name in drop:
        items.pop(name, None)

    # 消したパーツへの参照も取り除く（残すとファイルが壊れる）
    for rels in ('_rels/.rels', 'word/_rels/document.xml.rels'):
        xml = items[rels].decode()
        xml = re.sub(r'<Relationship\b[^>]*?'
                     r'(stylesWithEffects|thumbnail|customXml|numbering|theme|webSettings|'
                     r'extended-properties|settings|fontTable|core-properties)[^>]*?/>', '', xml)
        items[rels] = xml.encode()
    ct = items['[Content_Types].xml'].decode()
    for frag in ('stylesWithEffects', 'customXml', 'jpeg', 'numbering', 'theme',
                 'webSettings', 'extended-properties', 'settings', 'fontTable',
                 'core-properties'):
        ct = re.sub(r'<(?:Override|Default)\b[^>]*?' + frag + r'[^>]*?/>', '', ct)
    items['[Content_Types].xml'] = ct.encode()

    # styles.xml は雛形の定義がほとんど使われないので、必要な2つだけ残す
    items['word/styles.xml'] = MINIMAL_STYLES.encode()

    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as out:
        for name, data in items.items():
            out.writestr(name, data)


if __name__ == '__main__':
    to_name, issue_date, out = sys.argv[1], sys.argv[2], sys.argv[3]
    build(to_name, issue_date, out)
    print(f'{out}  {pathlib.Path(out).stat().st_size/1024:.1f} KB')
